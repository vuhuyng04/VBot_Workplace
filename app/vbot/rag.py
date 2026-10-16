"""Per-user retrieval: Postgres full-text + pgvector (Titan Embed v2, 512 dims) in one query.
Isolation: every query filters user_id AND runs under RLS (db.user_tx)."""
import io
import json
import re

from .aws import client
from .db import Tx

EMBED_MODEL = "amazon.titan-embed-text-v2:0"


def extract_text(name: str, data: bytes) -> str:
    n = name.lower()
    if n.endswith((".txt", ".md", ".csv", ".json", ".py", ".sql", ".html")):
        return data.decode("utf-8", errors="replace")
    if n.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    if n.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"# Sheet {ws.title}")
            for row in ws.iter_rows(values_only=True, max_row=2000):
                if any(v is not None for v in row):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if n.endswith(".pptx"):
        import pptx
        p = pptx.Presentation(io.BytesIO(data))
        return "\n".join(sh.text for s in p.slides for sh in s.shapes if getattr(sh, "has_text_frame", False))
    if n.endswith(".pdf"):
        import pypdf
        r = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((pg.extract_text() or "") for pg in r.pages[:200])
    return ""


def chunk(text: str, size: int = 1600, overlap: int = 200) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        i += size - overlap
    return out


def embed(texts: list[str]) -> list[list[float]] | None:
    """None when Bedrock is not reachable: search degrades to full-text only."""
    try:
        br = client("bedrock-runtime")
        vecs = []
        for t in texts:
            r = br.invoke_model(modelId=EMBED_MODEL, body=json.dumps({"inputText": t[:8000], "dimensions": 512, "normalize": True}))
            vecs.append(json.loads(r["body"].read())["embedding"])
        return vecs
    except Exception as e:
        print(json.dumps({"level": "warn", "msg": "embedding unavailable", "error": str(e)[:200]}))
        return None


def _vec(v):
    return "[" + ",".join(f"{x:.6f}" for x in v) + "]"


def index(tx: Tx, user_id: str, document_id: str, text: str) -> int:
    parts = chunk(text)
    vecs = embed(parts) if parts else None
    for i, p in enumerate(parts):
        if vecs:
            tx.run("INSERT INTO chunks (document_id, user_id, ord, content, embedding) VALUES (CAST(:d AS uuid), :u, :o, :c, CAST(:e AS halfvec))",
                   d=document_id, u=user_id, o=i, c=p, e=_vec(vecs[i]))
        else:
            tx.run("INSERT INTO chunks (document_id, user_id, ord, content) VALUES (CAST(:d AS uuid), :u, :o, :c)",
                   d=document_id, u=user_id, o=i, c=p)
    return len(parts)


def search(tx: Tx, user_id: str, query: str, k: int = 6) -> list[dict]:
    words = [w for w in re.findall(r"\w+", query.lower()) if len(w) > 1][:12]
    tsq = " | ".join(words) if words else ""
    vec = embed([query])
    if vec:
        rows = tx.run(
            """WITH v AS (SELECT c.id, 1 - (c.embedding <=> CAST(:e AS halfvec)) AS score FROM chunks c
                          WHERE c.user_id = :u AND c.embedding IS NOT NULL ORDER BY c.embedding <=> CAST(:e AS halfvec) LIMIT 20),
                    t AS (SELECT c.id, ts_rank(c.tsv, to_tsquery('simple', :q)) AS score FROM chunks c
                          WHERE c.user_id = :u AND :q <> '' AND c.tsv @@ to_tsquery('simple', :q) LIMIT 20)
               SELECT c.content, d.title, d.source, COALESCE(v.score, 0) * 0.7 + COALESCE(t.score, 0) * 0.3 AS score
               FROM chunks c JOIN documents d ON d.id = c.document_id
               LEFT JOIN v ON v.id = c.id LEFT JOIN t ON t.id = c.id
               WHERE c.user_id = :u AND (v.id IS NOT NULL OR t.id IS NOT NULL)
               ORDER BY score DESC LIMIT :k""",
            e=_vec(vec[0]), u=user_id, q=tsq, k=k)
    elif tsq:
        rows = tx.run(
            """SELECT c.content, d.title, d.source, ts_rank(c.tsv, to_tsquery('simple', :q)) AS score
               FROM chunks c JOIN documents d ON d.id = c.document_id
               WHERE c.user_id = :u AND c.tsv @@ to_tsquery('simple', :q)
               ORDER BY score DESC LIMIT :k""", u=user_id, q=tsq, k=k)
    else:
        rows = []
    return rows
